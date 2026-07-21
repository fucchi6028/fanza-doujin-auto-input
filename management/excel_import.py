"""
既存の「FANZA出品管理表」Excel を読み取って SQLite に取り込む。

想定レイアウト（実ファイルから確認）:
- シート = 月（"2026年4月" … "2027年3月"）
- 行2 = 見出し、行3〜 = アカウント行、末尾に「集計」行
- 列: A=アカウント(例 'アカウント001')  B=サークル名
      C/D/E=作品①②③（出品管理表, '○'=出品済み）  F=完了状況
      G/H/I=予告①②③（予告出品管理, '○'=出品済み）  J=予告完了状況

'○' を出品済み(チェック)とみなす。'×'・空白は未出品として扱う。
"""
import re

import openpyxl

from management import db
from management.db import (
    TABLE_LISTING, TABLE_RESERVED, SLOTS_PER_MONTH, normalize_account, now_iso,
)

# 列インデックス（1始まり）
COL_ACCOUNT = 1
COL_CIRCLE = 2
COLS_LISTING = (3, 4, 5)   # 作品①②③
COLS_RESERVED = (7, 8, 9)  # 予告①②③

CHECK_MARKS = {"○", "◯", "〇", "✓", "✔", "●", "☑"}


def _parse_ym(sheet_title: str):
    """'2026年4月' -> '2026-04'。解析できなければ None。"""
    m = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月", sheet_title)
    if not m:
        return None
    return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"


def _account_code(cell_value):
    """'アカウント001' -> '001' / 'アカウントK-01' -> 'K-01'。集計行等は None。"""
    if cell_value is None:
        return None
    s = str(cell_value).strip()
    if not s or s in ("集計", "合計"):
        return None
    s = re.sub(r"^アカウント", "", s).strip()
    if not s:
        return None
    return normalize_account(s)


def _is_checked(cell_value):
    if cell_value is None:
        return False
    s = str(cell_value).strip()
    if not s:
        return False
    return s in CHECK_MARKS


def import_workbook(manager, path, months=None, overwrite=True):
    """Excel を取り込む。

    overwrite=True: 各(月,アカウント,表)の枠を Excel の内容で置き換える。
    months 指定時はそのym集合のみ取り込む。

    戻り値: 取り込み結果サマリ dict。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    db_path = manager.db_path

    result = {
        "sheets": 0, "months": [], "accounts": set(),
        "listing_checks": 0, "reserved_checks": 0,
        "circle_names": 0, "skipped_sheets": [],
    }

    for ws in wb.worksheets:
        ym = _parse_ym(ws.title)
        if not ym or (months and ym not in months):
            result["skipped_sheets"].append(ws.title)
            continue
        result["sheets"] += 1
        result["months"].append(ym)

        # 対象月・全アカウントの枠を用意
        manager.ensure_months([ym])

        for r in range(3, ws.max_row + 1):
            code = _account_code(ws.cell(row=r, column=COL_ACCOUNT).value)
            if not code:
                continue
            result["accounts"].add(code)

            # サークル名（表示用アカウント名）
            circle = ws.cell(row=r, column=COL_CIRCLE).value
            if circle and str(circle).strip():
                db.set_account_name(code, str(circle).strip(), db_path)
                result["circle_names"] += 1
                db.upsert_account(code, 1, db_path)

            # 出品管理表 / 予告(予約)出品管理表 の各枠
            for table_type, cols, key in (
                (TABLE_LISTING, COLS_LISTING, "listing_checks"),
                (TABLE_RESERVED, COLS_RESERVED, "reserved_checks"),
            ):
                for idx, col in enumerate(cols, start=1):
                    checked = _is_checked(ws.cell(row=r, column=col).value)
                    if checked:
                        db.update_slot(ym, code, table_type, idx,
                                       checked=True, source_op="Excel取込",
                                       db_path=db_path)
                        db.add_log(ym, code, table_type, idx, "Excel取込",
                                   detail=f"{ws.title} から取込（枠{idx}）",
                                   db_path=db_path)
                        result[key] += 1
                    elif overwrite:
                        # Excelで空 = 未出品。既存チェックをクリア（取込由来のみ想定）
                        db.update_slot(ym, code, table_type, idx,
                                       checked=False, product_title="",
                                       db_path=db_path)

    result["accounts"] = sorted(result["accounts"], key=db.account_sort_key)
    return result

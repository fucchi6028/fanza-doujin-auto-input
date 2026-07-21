"""
出品管理表 / 予約出品管理表 の Excel エクスポート（閲覧・共有用）

- SQLite を正とし、都度この関数で読み取り専用の .xlsx を生成する。
- 自動化スクリプトは Excel を直接編集しない（ロック/競合回避）。
- 色分け: 3枠完了=緑 / 1〜2枠=黄 / 0枠=赤（アカウント行と消化数に付与）。
- 月ごとにシートを分け、過去月・当月・翌月を1ブックで参照できる。
"""
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

from management.db import TABLE_LISTING, TABLE_RESERVED, SLOTS_PER_MONTH, is_occupied, get_month_rows
from management import manager as mgr_mod

# Excel 標準の3色（条件付き書式と同じ配色）
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")
FILL_RED = PatternFill("solid", fgColor="FFC7CE")
FILL_HEADER = PatternFill("solid", fgColor="305496")
FONT_HEADER = Font(color="FFFFFF", bold=True)
FONT_GREEN = Font(color="006100")
FONT_YELLOW = Font(color="9C6500")
FONT_RED = Font(color="9C0006")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TABLE_SHEET_PREFIX = {TABLE_LISTING: "出品", TABLE_RESERVED: "予約"}


def _fill_for(count):
    if count >= SLOTS_PER_MONTH:
        return FILL_GREEN, FONT_GREEN, "完了"
    if count > 0:
        return FILL_YELLOW, FONT_YELLOW, "進行中"
    return FILL_RED, FONT_RED, "未着手"


def _cell_text(row):
    """枠セルの表示: チェック✓ + 作品名（+ステータス）。"""
    if not is_occupied(row):
        # 却下は空き扱いだが履歴として括弧付きで残す
        if row.get("status") == "却下" and row.get("product_title"):
            return f"（却下:{row['product_title']}）"
        return ""
    mark = "✓"
    parts = [mark]
    if row.get("product_title"):
        parts.append(row["product_title"])
    if row.get("status"):
        parts.append(f"[{row['status']}]")
    return " ".join(parts)


def _write_table_sheet(ws, manager, ym, table_type):
    accounts = manager.accounts()
    # スロットを (account -> {slot_no: row}) に整理
    by_acc = {a: {} for a in accounts}
    for r in get_month_rows(ym, table_type, manager.db_path):
        by_acc.setdefault(r["account"], {})[r["slot_no"]] = r

    names = manager.account_names()
    # 列: 1=アカウント 2=サークル名 3..=枠 その後 消化数/状態
    slot_base = 2  # 枠1 は列 slot_base+1 = 3
    cnt_col_idx = slot_base + SLOTS_PER_MONTH + 1
    st_col_idx = slot_base + SLOTS_PER_MONTH + 2

    headers = (["アカウント", "サークル名"]
               + [f"枠{i}" for i in range(1, SLOTS_PER_MONTH + 1)]
               + ["消化数", "状態"])
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = FILL_HEADER
        c.font = FONT_HEADER
        c.alignment = CENTER
        c.border = BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for i, account in enumerate(accounts, start=2):
        slots = by_acc.get(account, {})
        count = sum(1 for r in slots.values() if is_occupied(r))
        fill, font, label = _fill_for(count)

        acc_cell = ws.cell(row=i, column=1, value=account)
        acc_cell.fill = fill
        acc_cell.font = font
        acc_cell.alignment = CENTER
        acc_cell.border = BORDER

        circle_cell = ws.cell(row=i, column=2, value=names.get(account, ""))
        circle_cell.alignment = LEFT
        circle_cell.border = BORDER

        for s in range(1, SLOTS_PER_MONTH + 1):
            row = slots.get(s, {"checked": 0, "status": ""})
            cell = ws.cell(row=i, column=slot_base + s, value=_cell_text(row))
            cell.alignment = LEFT
            cell.border = BORDER

        cnt_cell = ws.cell(row=i, column=cnt_col_idx, value=count)
        cnt_cell.alignment = CENTER
        cnt_cell.fill = fill
        cnt_cell.font = font
        cnt_cell.border = BORDER

        st_cell = ws.cell(row=i, column=st_col_idx, value=label)
        st_cell.alignment = CENTER
        st_cell.fill = fill
        st_cell.font = font
        st_cell.border = BORDER

    # 列幅
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 16
    for s in range(1, SLOTS_PER_MONTH + 1):
        ws.column_dimensions[get_column_letter(slot_base + s)].width = 24
    ws.column_dimensions[get_column_letter(cnt_col_idx)].width = 8
    ws.column_dimensions[get_column_letter(st_col_idx)].width = 10

    # 消化数列に条件付き書式（0=赤 / 3=緑）も付与（Excel側で編集されても色が追随）
    last = len(accounts) + 1
    cnt_col = get_column_letter(cnt_col_idx)
    rng = f"{cnt_col}2:{cnt_col}{last}"
    ws.conditional_formatting.add(
        rng, CellIsRule(operator="equal", formula=["3"], fill=FILL_GREEN, font=FONT_GREEN))
    ws.conditional_formatting.add(
        rng, CellIsRule(operator="equal", formula=["0"], fill=FILL_RED, font=FONT_RED))
    ws.conditional_formatting.add(
        rng, CellIsRule(operator="between", formula=["1", "2"], fill=FILL_YELLOW, font=FONT_YELLOW))


def _write_dashboard(ws, manager, ym):
    ws["A1"] = f"ダッシュボード（{ym} 出品管理表）"
    ws["A1"].font = Font(bold=True, size=14)

    zero = manager.accounts_zero(ym, TABLE_LISTING)
    one = manager.accounts_one_left(ym, TABLE_LISTING)
    incomplete = manager.accounts_incomplete(ym, TABLE_LISTING)
    counts = manager.fill_counts(ym, TABLE_LISTING)
    done = [a for a, c in counts.items() if c >= SLOTS_PER_MONTH]

    rows = [
        ("完了（3/3）", len(done), "、".join(sorted(done, key=mgr_mod.account_sort_key))),
        ("3枠未満（未完了）", len(incomplete), "、".join(incomplete)),
        ("あと1枠だけ空き", len(one), "、".join(one)),
        ("未着手（0枠）", len(zero), "、".join(zero)),
    ]
    ws["A3"] = "区分"; ws["B3"] = "件数"; ws["C3"] = "アカウント"
    for c in ("A3", "B3", "C3"):
        ws[c].fill = FILL_HEADER; ws[c].font = FONT_HEADER; ws[c].alignment = CENTER
    for i, (label, n, names) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=n).alignment = CENTER
        ws.cell(row=i, column=3, value=names).alignment = LEFT

    # 出品ペース（日別消化枠数）
    pace = manager.pace(ym, TABLE_LISTING)
    ws["A9"] = "出品ペース（日別 消化枠数）"
    ws["A9"].font = Font(bold=True)
    ws["A10"] = "日付"; ws["B10"] = "消化枠数"
    for c in ("A10", "B10"):
        ws[c].fill = FILL_HEADER; ws[c].font = FONT_HEADER; ws[c].alignment = CENTER
    r = 11
    for day, n in pace.items():
        ws.cell(row=r, column=1, value=day)
        ws.cell(row=r, column=2, value=n).alignment = CENTER
        r += 1
    if not pace:
        ws.cell(row=11, column=1, value="（消化なし）")

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 80


def export_workbook(manager, path, months=None):
    """管理表ブックを生成する。months 未指定なら 先月・当月・翌月。"""
    cur = mgr_mod.current_ym()
    if months is None:
        months = [mgr_mod.prev_ym(cur), cur, mgr_mod.next_ym(cur)]

    wb = Workbook()
    wb.remove(wb.active)  # 既定シートを削除

    _write_dashboard(wb.create_sheet("ダッシュボード"), manager, cur)

    for ym in months:
        for table_type in (TABLE_LISTING, TABLE_RESERVED):
            name = f"{TABLE_SHEET_PREFIX[table_type]}_{ym}"
            _write_table_sheet(wb.create_sheet(name[:31]), manager, ym, table_type)

    wb.save(path)
    return path
